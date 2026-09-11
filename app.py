import io
import re
from datetime import date, datetime
import pandas as pd
import streamlit as st

st.set_page_config(page_title="WhatsApp Daily Report Automation", page_icon="📊", layout="wide")

COLS=["S.No","Ticket No","Start Date","End Date","Project Site","Activities","KE Representative","Vendor Supervisor","Manpower (Labour)"]
TEXT=["Ticket No","Project Site","Activities","KE Representative","Vendor Supervisor"]
DATES=["Start Date","End Date"]
DEDUP=["Ticket No","Start Date","End Date","Project Site","Activities","KE Representative","Vendor Supervisor","Manpower (Labour)"]

st.markdown("""<style>
.main-title{font-size:2.2rem;font-weight:700;margin-bottom:.2rem}.sub-title{color:#666;margin-bottom:1.5rem}.section-title{font-size:1.4rem;font-weight:700;margin:1.2rem 0 .8rem}.date-help{color:#555;font-size:.9rem;margin-bottom:8px}
</style>""",unsafe_allow_html=True)

def clean(v):
    if v is None:return ""
    try:
        if pd.isna(v):return ""
    except (TypeError,ValueError):pass
    s=str(v).strip()
    return "" if s.lower() in {"nan","nat","none"} else re.sub(r"\s+"," ",s)

def dt(v):
    if v is None:return pd.NaT
    try:
        if pd.isna(v):return pd.NaT
    except (TypeError,ValueError):pass
    if isinstance(v,pd.Timestamp):return v.normalize() if pd.notna(v) else pd.NaT
    if isinstance(v,datetime):return pd.Timestamp(v).normalize()
    if isinstance(v,date):return pd.Timestamp(v)
    if isinstance(v,(int,float)) and not pd.isna(v) and 20000<=float(v)<=60000:
        x=pd.to_datetime(v,unit="D",origin="1899-12-30",errors="coerce")
        return x.normalize() if pd.notna(x) else pd.NaT
    s=clean(v)
    if not s:return pd.NaT
    s=re.sub(r"(\d+)(st|nd|rd|th)\b",r"\1",s,flags=re.I)
    for f in ["%d.%m.%y","%d.%m.%Y","%d/%m/%y","%d/%m/%Y","%d-%m-%y","%d-%m-%Y","%d %B %Y","%d %b %Y","%d %B","%d %b"]:
        try:
            x=datetime.strptime(s,f)
            if "%y" not in f and "%Y" not in f:x=x.replace(year=date.today().year)
            return pd.Timestamp(x).normalize()
        except ValueError:pass
    try:
        x=pd.to_datetime(s,dayfirst=True,errors="coerce")
        return x.normalize() if pd.notna(x) else pd.NaT
    except Exception:return pd.NaT

def manpower(v):
    s=clean(v)
    if not s:return pd.NA
    m=re.search(r"[-+]?\d+(?:\.\d+)?",s)
    if not m:return pd.NA
    x=float(m.group())
    return int(x) if x.is_integer() else x

def norm(s):return re.sub(r"[\s_\-.]+"," ",clean(s).lower()).strip()

def standardize(df):
    df=df.copy().dropna(axis=0,how="all").dropna(axis=1,how="all")
    if df.empty:return pd.DataFrame(columns=COLS)
    n={norm(c):c for c in df.columns}
    aliases={
      "S.No":["s no","s.no","serial no","serial number","sr no"],
      "Ticket No":["ticket no","ticket number","ticket"],
      "Start Date":["start date","star date","start"],"End Date":["end date","end"],
      "Project Site":["project site","site","project"],"Activities":["activities","activity","work activity"],
      "KE Representative":["ke representative","ke supervisor","ke rep"],
      "Vendor Supervisor":["vendor supervisor","vendor sup","vendor representative","vendor"],
      "Manpower (Labour)":["manpower (labour)","manpower labour","manpower","labour","labor","manpower nos"]}
    out=pd.DataFrame(index=df.index)
    for c in COLS:
        src=None
        if c!="S.No":
            for a in aliases.get(c,[]):
                if norm(a) in n:src=n[norm(a)];break
        out[c]=df[src] if src is not None else pd.NA
    for c in TEXT:out[c]=out[c].apply(clean)
    for c in DATES:out[c]=out[c].apply(dt)
    out["Manpower (Labour)"]=out["Manpower (Labour)"].apply(manpower)
    return out[COLS]

def read_excel(f):
    raw=pd.read_excel(io.BytesIO(f.getvalue()),sheet_name=0,header=None)
    if raw.empty:return pd.DataFrame(columns=COLS),"Excel file is empty."
    h=None
    for i in range(min(len(raw),30)):
        if any(norm(x) in {"ticket no","ticket number","ticket"} for x in raw.iloc[i]):h=i;break
    try:df=pd.read_excel(io.BytesIO(f.getvalue()),sheet_name=0,header=0 if h is None else h)
    except Exception as e:return pd.DataFrame(columns=COLS),str(e)
    out=standardize(df);out=out[out["Ticket No"].apply(clean)!=""].copy()
    return out,None

def blocks(text):
    text=text.replace("\r\n","\n").replace("\r","\n").strip()
    ms=list(re.finditer(r"(?im)^\s*\*?\s*Ticket\s*No\.?\s*[:\-]?\s*",text))
    return [text[m.start():(ms[i+1].start() if i+1<len(ms) else len(text))].strip() for i,m in enumerate(ms)] if ms else []

def field(b,patterns):
    for p in patterns:
        m=re.search(p,b,re.I|re.M)
        if m:return clean(m.group(1))
    return ""

def parse(text):
    rows=[]
    for b in blocks(text):
        ke=field(b,[r"^\s*\*?\s*KE\s*Supervisor\s*[:\-]?\s*(.+?)\s*$",r"^\s*\*?\s*KE\s*Representative\s*[:\-]?\s*(.+?)\s*$"])
        if ke.lower() in {"yes","no","y","n","ok","available"}:ke=""
        rows.append({
          "S.No":pd.NA,
          "Ticket No":field(b,[r"^\s*\*?\s*Ticket\s*No\.?\s*[:\-]?\s*(.+?)\s*$"]),
          "Start Date":dt(field(b,[r"^\s*\*?\s*Start\s*Date\s*[:\-]?\s*(.+?)\s*$",r"^\s*\*?\s*Star\s*Date\s*[:\-]?\s*(.+?)\s*$"])),
          "End Date":dt(field(b,[r"^\s*\*?\s*End\s*Date\s*[:\-]?\s*(.+?)\s*$"])),
          "Project Site":field(b,[r"^\s*\*?\s*Project\s*Site\s*[:\-]?\s*(.+?)\s*$",r"^\s*\*?\s*Project\s*[:\-]?\s*(.+?)\s*$"]),
          "Activities":field(b,[r"^\s*\*?\s*Activity\s*[:\-]?\s*(.+?)\s*$",r"^\s*\*?\s*Activities\s*[:\-]?\s*(.+?)\s*$"]),
          "KE Representative":ke,
          "Vendor Supervisor":field(b,[r"^\s*\*?\s*Vendor\s*Supervisor\s*[:\-]?\s*(.+?)\s*$",r"^\s*\*?\s*Vendor\s*[:\-]?\s*(.+?)\s*$"]),
          "Manpower (Labour)":manpower(field(b,[r"^\s*\*?\s*Manpower(?:\s*\(Labour\))?\s*[:\-]?\s*(.+?)\s*$"]))})
    return standardize(pd.DataFrame(rows,columns=COLS))

def prepare(df):
    x=standardize(df)
    x["S.No"]=range(1,len(x)+1)
    for c in DATES:x[c]=pd.to_datetime(x[c],errors="coerce")
    x["Manpower (Labour)"]=pd.to_numeric(x["Manpower (Labour)"],errors="coerce").astype("Float64")
    return x

def overlap(df,a,b):
    x=prepare(df)
    if x.empty:return x
    ae=x["End Date"].fillna(x["Start Date"])
    return x[x["Start Date"].notna() & ae.notna() & (x["Start Date"]<=pd.Timestamp(b)) & (ae>=pd.Timestamp(a))].copy()

def validate(df):
    req=["Ticket No","Start Date","End Date","Project Site","Activities","Manpower (Labour)"]
    return pd.Series(["Incomplete" if any(pd.isna(r[c]) or (isinstance(r[c],str) and not r[c].strip()) for c in req) else "Complete" for _,r in df.iterrows()],index=df.index)

def export(df,a,b):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    x=prepare(df);buf=io.BytesIO();status=validate(x)
    with pd.ExcelWriter(buf,engine="openpyxl",date_format="DD-MM-YYYY",datetime_format="DD-MM-YYYY") as w:
        x.to_excel(w,sheet_name="Daily Report",index=False,startrow=3)
        pd.DataFrame({"Metric":["Report From","Report To","Total Tickets","Complete Records","Incomplete Records","Total Manpower"],"Value":[a,b,len(x),int((status=="Complete").sum()),int((status=="Incomplete").sum()),float(pd.to_numeric(x["Manpower (Labour)"],errors="coerce").fillna(0).sum())]}).to_excel(w,sheet_name="Summary",index=False)
    buf.seek(0);wb=load_workbook(buf);ws=wb["Daily Report"];ws["A1"]="Daily Site Manpower & Work Progress Summary";ws["A2"]=f"Report Period: {pd.Timestamp(a):%d-%b-%Y} to {pd.Timestamp(b):%d-%b-%Y}";ws["A1"].font=Font(bold=True,size=16);ws["A2"].font=Font(italic=True)
    for c in ws[4]:c.fill=PatternFill("solid",fgColor="1F4E78");c.font=Font(bold=True,color="FFFFFF");c.alignment=Alignment(horizontal="center",wrap_text=True)
    for r in range(5,ws.max_row+1):
        ws.cell(r,3).number_format="DD-MM-YYYY";ws.cell(r,4).number_format="DD-MM-YYYY"
        for c in ws[r]:c.alignment=Alignment(vertical="top",wrap_text=True)
    for col,wid in zip("ABCDEFGHI",[8,16,14,14,28,42,24,24,18]):ws.column_dimensions[col].width=wid
    ws.freeze_panes="A5";ws.auto_filter.ref=ws.dimensions
    out=io.BytesIO();wb.save(out);out.seek(0);return out

for k in ["excel_df","txt_df","paste_df","filtered_df"]:
    if k not in st.session_state:st.session_state[k]=pd.DataFrame(columns=COLS)
for k in ["excel_name","txt_name"]:
    if k not in st.session_state:st.session_state[k]=""

st.markdown('<div class="main-title">📊 WhatsApp Daily Report Automation</div>',unsafe_allow_html=True)
st.markdown('<div class="sub-title">Merge Excel records with WhatsApp TXT and pasted messages, filter by activity period, review, and download a consolidated Excel.</div>',unsafe_allow_html=True)

st.markdown('<div class="section-title">1. Select Report Period</div>',unsafe_allow_html=True)
c1,c2=st.columns(2)
with c1:a=st.date_input("From Date",value=date.today(),format="DD/MM/YYYY",key="from_date")
with c2:b=st.date_input("To Date",value=date.today(),format="DD/MM/YYYY",key="to_date")
if a>b:st.error("From Date cannot be later than To Date.");st.stop()
st.info(f"Selected period: **{a:%d-%b-%Y}** to **{b:%d-%b-%Y}**")

st.markdown('<div class="section-title">2. Upload Existing Excel Report</div>',unsafe_allow_html=True)
ef=st.file_uploader("Upload your existing report Excel",type=["xlsx","xls","xlsm"],key="excel_upload")
if ef is not None and ef.name!=st.session_state.excel_name:
    x,e=read_excel(ef)
    if e:st.error(f"Excel reading error: {e}")
    else:st.session_state.excel_df=x;st.session_state.excel_name=ef.name;st.success(f"Excel loaded: **{len(x)} records**.")

st.markdown('<div class="section-title">3. Add WhatsApp Report Data</div>',unsafe_allow_html=True)
tf=st.file_uploader("A. Upload WhatsApp TXT export",type=["txt"],key="txt_upload")
if tf is not None and tf.name!=st.session_state.txt_name:
    try:
        x=parse(tf.getvalue().decode("utf-8",errors="ignore"));st.session_state.txt_df=x;st.session_state.txt_name=tf.name;st.success(f"TXT WhatsApp source parsed: **{len(x)} records**.")
    except Exception as e:st.error(f"WhatsApp TXT parsing error: {e}")

p=st.text_area("B. Paste additional WhatsApp messages",height=240,placeholder="* Ticket No.298662\n* Start Date 9.9.26\n* End Date 9.9.26\n* Project Site Queens road\n* Activity Paint work\n* KE Supervisor Zulfiqar\n* Vendor Supervisor Abdullah\n* Manpower 2nos",key="paste_text")
if st.button("Parse Pasted Messages",type="primary"):
    if not p.strip():st.warning("Please paste WhatsApp messages first.")
    else:
        x=parse(p);st.session_state.paste_df=x;st.success(f"Pasted WhatsApp source parsed: **{len(x)} records**.")

tc,pc=len(st.session_state.txt_df),len(st.session_state.paste_df)
st.success(f"WhatsApp sources ready: **{tc} TXT + {pc} pasted = {tc+pc} records before duplicate removal.**")

st.markdown('<div class="section-title">4. Merge & Filter Records</div>',unsafe_allow_html=True)
if st.button("🔄 Merge Excel + All WhatsApp Data",type="primary",use_container_width=True):
    parts=[x for x in [st.session_state.excel_df,st.session_state.txt_df,st.session_state.paste_df] if not x.empty]
    if not parts:st.error("Please upload Excel or add WhatsApp data first.")
    else:
        m=standardize(pd.concat(parts,ignore_index=True));m=m[m["Ticket No"].apply(clean)!=""].copy();before=len(m);m=m.drop_duplicates(subset=DEDUP,keep="first");dups=before-len(m);f=overlap(m,a,b);st.session_state.filtered_df=f;st.success(f"Merge completed: **{len(m)} unique records**; **{dups} exact duplicate(s)** removed; **{len(f)} record(s)** overlap the selected period.")

st.markdown('<div class="section-title">5. Review & Edit Report</div>',unsafe_allow_html=True)
if st.session_state.filtered_df.empty:st.info("No filtered records yet. Upload data and click **Merge Excel + All WhatsApp Data**.")
else:
    ed=prepare(st.session_state.filtered_df)
    ed=st.data_editor(ed,column_config={"S.No":st.column_config.NumberColumn("S.No",disabled=True),"Ticket No":st.column_config.TextColumn("Ticket No"),"Start Date":st.column_config.DateColumn("Start Date",format="DD/MM/YYYY"),"End Date":st.column_config.DateColumn("End Date",format="DD/MM/YYYY"),"Project Site":st.column_config.TextColumn("Project Site"),"Activities":st.column_config.TextColumn("Activities"),"KE Representative":st.column_config.TextColumn("KE Representative"),"Vendor Supervisor":st.column_config.TextColumn("Vendor Supervisor"),"Manpower (Labour)":st.column_config.NumberColumn("Manpower (Labour)",min_value=0,step=1)},hide_index=True,use_container_width=True,num_rows="dynamic",key="editor")
    st.session_state.filtered_df=prepare(ed)

if not st.session_state.filtered_df.empty:
    x=prepare(st.session_state.filtered_df);s=validate(x);tot=len(x);comp=int((s=="Complete").sum());inc=int((s=="Incomplete").sum());mp=pd.to_numeric(x["Manpower (Labour)"],errors="coerce").fillna(0).sum()
    st.markdown('<div class="section-title">6. Report Summary</div>',unsafe_allow_html=True)
    q1,q2,q3,q4=st.columns(4);q1.metric("Total Records",tot);q2.metric("Complete",comp);q3.metric("Incomplete",inc);q4.metric("Total Manpower",int(mp) if float(mp).is_integer() else round(float(mp),2))
    if inc:st.warning(f"{inc} record(s) have missing required fields.")
    else:st.success("All displayed records contain the required fields.")
    st.markdown('<div class="section-title">7. Download Consolidated Report</div>',unsafe_allow_html=True)
    buf=export(x,a,b);st.download_button("📥 Download Consolidated Excel",data=buf.getvalue(),file_name=f"Merged_Daily_Report_{a:%d-%m-%Y}_to_{b:%d-%m-%Y}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",type="primary",use_container_width=True)

with st.sidebar:
    st.header("Instructions")
    st.markdown("""1. Select From/To dates.  \n2. Upload Excel.  \n3. Upload TXT and/or paste additional WhatsApp messages.  \n4. Parse pasted messages.  \n5. Merge all sources.  \n6. Review/edit and download.  \n\n**Date filter:** activity-period overlap is used.  \n**Duplicates:** Ticket No alone is not used; exact duplicates use all report fields except S.No.\n\nKeep operational Excel/WhatsApp data out of GitHub.""")
